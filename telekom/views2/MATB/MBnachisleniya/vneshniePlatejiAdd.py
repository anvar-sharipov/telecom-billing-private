from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import HttpResponse
from django.db import transaction

from datetime import datetime
import io
import re
import urllib.parse
import openpyxl
import tablib

from telekom.models import UserTable, OldLoginDogowor, PlatejiWhichAddKassirsEveryDay, SaveInfoAboutWhoAddAndNachPaysFromBilling, StaffAction

import logging

logger = logging.getLogger(__name__)


# Milli Billing "Online" источники внешних платежей — доверенный список.
# Всё что приходит из этого файла — 100% внешние платежи (kassir_etrap='Внешние платежи').
KASSA_MANAGERS = ['Capar', 'eGov', 'Saray', 'Toleg', 'Turkmenpost diller']

TARIFF_TO_TYPE = {'Internet': 'Internet', 'Älem TV': 'Alem', 'Telefon': 'Abonplata'}

ETRAP_CODES = {
    'Dashoguz': '322',
    'Akdepe': '344',
    'Boldumsaz': '346',
    'Gorogly': '340',
    'Koneurgench': '347',
    'Turkmenbashy': '349',
    'S.A.Nyyazow': '348',
    'Ruhubelent': '342',
    'Garashsyzlyk': '343',
    'Gubadag': '345',
}
CODE_TO_ETRAP = {v: k for k, v in ETRAP_CODES.items()}

EXPECTED_HEADER_CELLS = {
    0: 'Operator',
    4: 'Nyrhnama Topary',
    5: 'Töleg Senesi',
    8: 'Töleg Belgisi',
    10: 'Şertnama Belgisi',
    15: 'Müşderiniň Doly Ady',
    23: 'Möçberi',
}

months_ru = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь', 'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']


def build_problem_rows_xlsx_response(rows):
    headers = ('KASSA', 'Услуга', 'Пользователь', 'Договор', '№ платежа', 'Дата платежа', 'Сумма', 'Причина')
    data = tablib.Dataset(headers=headers)
    for row in rows:
        data.append((row['manager'], row['tariff'], row['fio'], row['dogowor'], row['pay_id'], row['date'], row['amount'], row['reason']))

    filename = f"vneshniePlateji_problem_rows_{datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx"
    encoded_filename = urllib.parse.quote(filename.encode('utf-8'))

    response = HttpResponse(data.xlsx, content_type='application/vnd.ms-excel;charset=utf-8')
    response['Content-Disposition'] = f"attachment; filename= {encoded_filename}"
    return response


def build_lookup_dicts():
    """Глобальные (по всем этрапам) словари для сопоставления договора с абонентом."""
    dogowor_to_user = {}
    dogowor_alem_to_user = {}
    dogowor_telefoniya_to_user = {}
    user_by_number_etrap = {}
    for u in UserTable.objects.all():
        user_by_number_etrap[(u.number, u.etrap)] = u
        if u.dogowor:
            dogowor_to_user.setdefault(u.dogowor.strip().upper(), u)
        if u.dogowor_alem:
            dogowor_alem_to_user.setdefault(u.dogowor_alem.strip().upper(), u)
        if u.dogowor_telefoniya:
            dogowor_telefoniya_to_user.setdefault(u.dogowor_telefoniya.strip().upper(), u)

    old_dogowor = {}
    old_dogowor_alem = {}
    old_dogowor_telefoniya = {}
    for o in OldLoginDogowor.objects.all():
        if o.dogowor:
            old_dogowor.setdefault(o.dogowor.strip().upper(), o)
        if o.dogowor_alem:
            old_dogowor_alem.setdefault(o.dogowor_alem.strip().upper(), o)
        if o.dogowor_telefoniya:
            old_dogowor_telefoniya.setdefault(o.dogowor_telefoniya.strip().upper(), o)

    return {
        'dogowor_to_user': dogowor_to_user,
        'dogowor_alem_to_user': dogowor_alem_to_user,
        'dogowor_telefoniya_to_user': dogowor_telefoniya_to_user,
        'user_by_number_etrap': user_by_number_etrap,
        'old_dogowor': old_dogowor,
        'old_dogowor_alem': old_dogowor_alem,
        'old_dogowor_telefoniya': old_dogowor_telefoniya,
    }


def _resolve_by_code(dogowor_raw, lookups):
    digits = re.sub(r'\D', '', dogowor_raw)
    if len(digits) not in (10, 11):
        return None
    code = digits[-8:-5]
    number = digits[-5:]
    etrap = CODE_TO_ETRAP.get(code)
    if not etrap:
        return None
    return lookups['user_by_number_etrap'].get((number, etrap))


def resolve_subscriber(tariff, dogowor_raw, lookups):
    """Возвращает (number, etrap, is_enterprises) или None."""
    key = dogowor_raw.strip().upper()

    found = None
    if tariff == 'Internet':
        found = lookups['dogowor_to_user'].get(key) or lookups['old_dogowor'].get(key)
    elif tariff == 'Älem TV':
        found = _resolve_by_code(dogowor_raw, lookups)
        if not found:
            found = lookups['dogowor_alem_to_user'].get(key) or lookups['old_dogowor_alem'].get(key)
    elif tariff == 'Telefon':
        found = _resolve_by_code(dogowor_raw, lookups)
        if not found:
            found = lookups['dogowor_telefoniya_to_user'].get(key) or lookups['old_dogowor_telefoniya'].get(key)

    if not found:
        return None
    return found.number, found.etrap, str(found.is_enterprises)


def parse_workbook(uploaded_file):
    """Читает grouped-отчёт milli billing 'Online toleg jemi'. Возвращает (rows, header_error)."""
    wb = openpyxl.load_workbook(io.BytesIO(uploaded_file.read()), data_only=True)
    ws = wb.worksheets[0]
    all_rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    header_row = None
    for r in all_rows[:10]:
        if r and r[0] and str(r[0]).strip() == 'Operator':
            header_row = r
            break

    if header_row is None:
        return None, 'Не найдена строка заголовков ("Operator") в первых 10 строках файла — формат файла не распознан.'

    for idx, expected in EXPECTED_HEADER_CELLS.items():
        actual = str(header_row[idx]).strip() if header_row[idx] is not None else ''
        if actual != expected:
            return None, f'Формат файла изменился: в колонке {idx} ожидалось "{expected}", получено "{actual}".'

    rows = []
    current_manager = None
    for r in all_rows:
        marker = str(r[0]).strip() if r[0] else ''
        if marker == 'KASSA':
            current_manager = str(r[1]).strip() if r[1] else ''
            continue
        if marker == 'Online':
            rows.append({
                'manager': current_manager,
                'tariff': r[4],
                'date': r[5],
                'pay_id': r[8],
                'dogowor': str(r[10]).strip() if r[10] is not None else '',
                'fio': str(r[15]).strip() if r[15] is not None else '',
                'amount': r[23],
            })

    return rows, None


def vneshniePlatejiAdd(request):
    context = {}

    if not request.user.is_authenticated:
        messages.error(request, 'Вы не аутентифицированы')
        return redirect('user-login')

    if not (request.user.is_superuser or request.user.username in ['Gayyp', 'yhlas_mtb']):
        messages.error(request, 'Доступ разрешен только администратору')
        return redirect('user-login')

    context['vneshniePlatejiAdd'] = True
    context['matbIndex'] = True
    context['months'] = months_ru
    context['years'] = ['2025', '2026', '2027', '2028']

    year = request.POST.get('year') or ''
    month = request.POST.get('month') or ''
    context['year'] = year
    context['month'] = month

    if request.method == 'POST':
        if not year or not month or month not in months_ru:
            messages.error(request, 'Выберите год и месяц')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        month_num = months_ru.index(month) + 1

        try:
            uploaded_file = request.FILES['my_file']
        except Exception:
            messages.error(request, 'Выберите файл')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        if not str(uploaded_file).lower().endswith('.xlsx'):
            messages.error(request, 'Файл должен быть формата xlsx')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        if PlatejiWhichAddKassirsEveryDay.objects.filter(file_name=str(uploaded_file)).exists():
            messages.error(request, f'Файл "{str(uploaded_file)}" уже был добавлен ранее.')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        try:
            parsed_rows, header_error = parse_workbook(uploaded_file)
        except Exception as e:
            messages.error(request, f'Ошибка чтения файла: {e}')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        if header_error:
            messages.error(request, header_error)
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        if not parsed_rows:
            messages.error(request, 'В файле не найдено ни одной строки платежей')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        # Защита от дурака: даты платежей должны попадать в выбранный год/месяц
        bad_dates = sorted({r['date'].strftime('%d.%m.%Y') for r in parsed_rows if not r['date'] or r['date'].year != int(year) or r['date'].month != month_num})
        if bad_dates:
            messages.error(request, f'В файле есть платежи с датой вне выбранного периода {month} {year}: {", ".join(bad_dates[:10])}{"..." if len(bad_dates) > 10 else ""}')
            return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)

        lookups = build_lookup_dicts()

        existing_pay_ids = set(
            PlatejiWhichAddKassirsEveryDay.objects.filter(manager__in=KASSA_MANAGERS).exclude(kodOplaty='').exclude(kodOplaty__isnull=True).values_list('kodOplaty', flat=True)
        )
        seen_pay_ids_in_file = set()

        rows_to_create = []
        problem_rows = []

        for r in parsed_rows:
            pay_id_str = str(r['pay_id']) if r['pay_id'] is not None else ''

            if not r['manager'] or r['manager'] not in KASSA_MANAGERS:
                problem_rows.append({**r, 'reason': f'Неизвестная KASSA-группа "{r["manager"]}" (не входит в список {KASSA_MANAGERS})'})
                continue

            if r['tariff'] not in TARIFF_TO_TYPE:
                problem_rows.append({**r, 'reason': f'Неизвестный тип услуги "{r["tariff"]}"'})
                continue

            if not pay_id_str:
                problem_rows.append({**r, 'reason': 'Пустой № платежа (Töleg Belgisi)'})
                continue

            if pay_id_str in existing_pay_ids:
                problem_rows.append({**r, 'reason': f'Платёж № {pay_id_str} уже был добавлен ранее (дубликат)'})
                continue

            if pay_id_str in seen_pay_ids_in_file:
                problem_rows.append({**r, 'reason': f'Платёж № {pay_id_str} повторяется несколько раз в этом же файле'})
                continue
            seen_pay_ids_in_file.add(pay_id_str)

            try:
                amount = float(r['amount'])
            except (TypeError, ValueError):
                problem_rows.append({**r, 'reason': f'Некорректная сумма "{r["amount"]}"'})
                continue

            resolved = resolve_subscriber(r['tariff'], r['dogowor'], lookups)
            if not resolved:
                problem_rows.append({**r, 'reason': 'Абонент не найден в UserTable/OldLoginDogowor'})
                continue

            number, etrap, is_enterprises = resolved

            rows_to_create.append(PlatejiWhichAddKassirsEveryDay(
                number=number,
                user_etrap=etrap,
                kassir_etrap='Внешние платежи',
                type_pay=TARIFF_TO_TYPE[r['tariff']],
                pay_category='Внешние Платежи',
                manager=r['manager'],
                date=r['date'].date(),
                kodOplaty=pay_id_str,
                dogowor=r['dogowor'],
                name=r['fio'],
                is_enterprises=is_enterprises,
                price=amount,
                file_name=str(uploaded_file),
                who_add_file=request.user.username,
            ))

        if request.POST.get('vneshniePlatejiExportProblems'):
            if not problem_rows:
                messages.success(request, 'Проблемных строк нет')
                return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)
            return build_problem_rows_xlsx_response(problem_rows)

        totalCount = len(parsed_rows)
        context['totalCount'] = totalCount
        context['matched_count'] = len(rows_to_create)
        context['problem_count'] = len(problem_rows)
        context['totalAmount'] = float('%.2f' % sum(r.price for r in rows_to_create))
        context['is_test_mode'] = bool(request.POST.get('vneshniePlatejiTest'))

        if context['is_test_mode']:
            messages.success(request, 'Проверка прошла успешно (ничего не сохранено в БД)')
        elif problem_rows:
            messages.error(request, f'Добавление запрещено: {len(problem_rows)} проблемных строк из {totalCount}. Скачайте список проблемных строк, исправьте и загрузите файл заново.')
        else:
            try:
                with transaction.atomic():
                    if PlatejiWhichAddKassirsEveryDay.objects.select_for_update().filter(file_name=str(uploaded_file)).exists():
                        messages.error(request, f'Файл "{str(uploaded_file)}" уже был добавлен ранее (повторная проверка).')
                        return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)
                    PlatejiWhichAddKassirsEveryDay.objects.bulk_create(rows_to_create)
                    SaveInfoAboutWhoAddAndNachPaysFromBilling.objects.create(
                        file_name=str(uploaded_file),
                        etrap_add='Внешние платежи',
                        who_add=request.user.username,
                        when_add=datetime.now(),
                    )
                    StaffAction.objects.create(
                        user=request.user,
                        comment=f'Добавление в PlatejiWhichAddKassirsEveryDay (Внешние платежи Milli Billing), файл {str(uploaded_file)}, строк {len(rows_to_create)}, добавил {request.user.username}',
                        action='Импорт с xlsx внешние платежи в БД',
                    )
                messages.success(request, f'Успешно добавлено {len(rows_to_create)} записей')
            except Exception as e:
                messages.error(request, f'Ошибка при сохранении, тип ошибки == {e}')
                logger.error(f'Ошибка при сохранении PlatejiWhichAddKassirsEveryDay (внешние платежи), тип ошибки == {e}')

    return render(request, 'telekom/MATB/MBnachisleniya/vneshniePlatejiAdd.html', context)
